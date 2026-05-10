from __future__ import annotations

import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from ppt_ui.core.component import RenderContext
from ppt_ui.core.layout import Box


def _items(value: object) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    return [dict(item) for item in value if isinstance(item, Mapping)]


def _style(ctx: RenderContext, key: str, default: str) -> str:
    value = ctx.style.get(key)
    return str(value) if value not in (None, "") else default


def _style_num(ctx: RenderContext, key: str, default: float) -> float:
    value = ctx.style.get(key)
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _color(ctx: RenderContext, value: object, default: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        raw = default
    colors = ctx.theme.colors
    resolved = getattr(colors, raw, raw)
    if re.match(r"^[0-9A-Fa-f]{6}$", str(resolved)):
        return f"#{resolved}"
    return str(resolved)


def _palette(ctx: RenderContext, fallback: list[str] | None = None) -> list[str]:
    palette = [str(color).strip().lstrip("#") for color in getattr(ctx.theme, "chart_palette", []) if str(color).strip()]
    return palette or fallback or [ctx.theme.colors.primary, ctx.theme.colors.accent, ctx.theme.colors.success, ctx.theme.colors.warning]


def _numbers(value: object) -> list[float]:
    if not isinstance(value, list):
        return []
    result: list[float] = []
    for item in value:
        try:
            result.append(float(item))
        except (TypeError, ValueError):
            result.append(0.0)
    return result


@dataclass
class TextComponent:
    text: str = ""
    bullets: list[str] = field(default_factory=list)
    size: int | None = None
    color: str | None = None
    bold: bool = False
    align: str = "left"
    valign: str = "top"

    def render(self, ctx: RenderContext, box: Box) -> None:
        if ctx.style.get("fill") or ctx.style.get("border"):
            ctx.renderer.add_card(
                ctx.slide,
                box,
                fill=_style(ctx, "fill", ctx.theme.colors.surface_white),
                line=_style(ctx, "border", ctx.theme.colors.border_light),
                shadow=bool(ctx.style.get("shadow", False)),
            )
            box = box.inset(0.18, 0.12)
        size = int(ctx.style.get("size", self.size or ctx.theme.fonts.body_size))
        color = _style(ctx, "color", self.color or ctx.theme.colors.text_secondary)
        bold = bool(ctx.style.get("bold", self.bold))
        align = str(ctx.style.get("align", self.align))
        valign = str(ctx.style.get("valign", self.valign))
        if self.bullets:
            ctx.renderer.bullet_list(ctx.slide, box, self.bullets, size=size, text_color=color, bullet_color=_style(ctx, "accent", ctx.theme.colors.primary))
            return
        ctx.renderer.text(
            ctx.slide,
            box,
            self.text,
            size=size,
            color=color,
            bold=bold,
            align=align,
            valign=valign,
        )

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "TextComponent":
        return cls(
            text=str(props.get("text", "")),
            bullets=[str(item) for item in props.get("bullets", [])],
            size=int(props["size"]) if props.get("size") is not None else None,
            color=str(props["color"]) if props.get("color") is not None else None,
            bold=bool(props.get("bold", False)),
            align=str(props.get("align", "left")),
            valign=str(props.get("valign", "top")),
        )


@dataclass
class MetricCardsComponent:
    cards: list[dict[str, Any]] = field(default_factory=list)

    def render(self, ctx: RenderContext, box: Box) -> None:
        count = max(1, len(self.cards))
        columns = int(ctx.style.get("columns", 0) or 0)
        if columns <= 0:
            columns = min(count, 4)
            cell_w = (box.w - ctx.theme.spacing.gutter * (columns - 1)) / max(1, columns)
            if cell_w < 1.55 and columns > 2:
                columns = 2
        columns = max(1, min(columns, count))
        rows_count = max(1, (count + columns - 1) // columns)
        rows = box.split_rows(rows_count, ctx.theme.spacing.gutter)
        cells = [cell for row in rows for cell in row.split_cols(columns, ctx.theme.spacing.gutter)]
        for cell, card in zip(cells, self.cards):
            ctx.renderer.add_metric_card(
                ctx.slide,
                cell,
                label=str(card.get("label", "")),
                value=str(card.get("value", "")),
                delta=str(card.get("delta", "")),
                note=str(card.get("compare", card.get("note", ""))),
                icon=str(card.get("icon", ""))[:3].upper(),
                fill=_style(ctx, "fill", ctx.theme.component_style("metric_card").fill),
                border=_style(ctx, "border", ctx.theme.component_style("metric_card").border),
                accent=_style(ctx, "accent", ctx.theme.component_style("metric_card").accent),
            )

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "MetricCardsComponent":
        return cls(cards=_items(props.get("cards", [])))


@dataclass
class LineChartComponent:
    categories: list[str] = field(default_factory=list)
    series: list[dict[str, Any]] = field(default_factory=list)

    def render(self, ctx: RenderContext, box: Box) -> None:
        r = ctx.renderer
        r.add_card(ctx.slide, box, fill=_style(ctx, "fill", ctx.theme.colors.surface_white), line=_style(ctx, "border", ctx.theme.colors.border_light))
        chart = box.inset(0.34, 0.34, 0.28, 0.38)
        legend_h = 0.24 if self.series else 0.0
        plot = Box(chart.x + 0.45, chart.y + 0.25 + legend_h, chart.w - 0.65, chart.h - 0.78 - legend_h)
        values = [float(value) for item in self.series for value in item.get("values", [])]
        max_value = max(values) if values else 1
        min_value = min(0, min(values) if values else 0)
        span = max(max_value - min_value, 1)
        caption_size = ctx.theme.fonts.caption_size
        line_width = _style_num(ctx, "line_width", 1.3)

        for idx in range(5):
            y = plot.y + idx * plot.h / 4
            r.line(ctx.slide, plot.x, y, plot.x + plot.w, y, ctx.theme.colors.border_light, width=0.6)
            label = int(max_value - idx * span / 4)
            r.text(ctx.slide, Box(chart.x, y - 0.06, 0.34, 0.12), str(label), size=caption_size, color=ctx.theme.colors.text_tertiary, align="right")

        colors = _palette(ctx)
        category_count = max(1, len(self.categories) - 1)
        for series_idx, item in enumerate(self.series):
            color = str(item.get("color", "")) or colors[series_idx % len(colors)]
            points: list[tuple[float, float]] = []
            for idx, raw_value in enumerate(item.get("values", [])):
                value = float(raw_value)
                x = plot.x + idx * plot.w / category_count
                y = plot.y + plot.h - ((value - min_value) / span) * plot.h
                points.append((x, y))
                r.circle(ctx.slide, Box(x - 0.035, y - 0.035, 0.07, 0.07), color, line="FFFFFF")
            for start, end in zip(points, points[1:]):
                r.line(ctx.slide, start[0], start[1], end[0], end[1], color, width=line_width)

        for idx, label in enumerate(self.categories):
            x = plot.x + idx * plot.w / category_count
            r.text(ctx.slide, Box(x - 0.22, plot.y + plot.h + 0.12, 0.44, 0.12), label, size=caption_size, color=ctx.theme.colors.text_secondary, align="center")

        legend_width = min(chart.w - 0.72, max(1.20, len(self.series[:4]) * 1.32))
        legend_x = chart.x + chart.w - legend_width
        cursor_x = legend_x
        for idx, item in enumerate(self.series[:4]):
            color = str(item.get("color", "")) or colors[idx % len(colors)]
            name = str(item.get("name", ""))
            item_width = min(1.55, max(1.05, 0.36 + len(name) * 0.055))
            r.rect(ctx.slide, Box(cursor_x, chart.y + 0.07, 0.13, 0.04), color, rounded=True)
            r.text(ctx.slide, Box(cursor_x + 0.18, chart.y + 0.005, item_width - 0.18, 0.16), name, size=caption_size, color=ctx.theme.colors.text_secondary)
            cursor_x += item_width + 0.14

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "LineChartComponent":
        return cls(
            categories=[str(item) for item in props.get("categories", [])],
            series=_items(props.get("series", [])),
        )


@dataclass
class ComparisonTableComponent:
    headers: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    conclusion: str = ""

    def render(self, ctx: RenderContext, box: Box) -> None:
        table_box = box
        if self.conclusion:
            table_box = Box(box.x, box.y, box.w, max(0.2, box.h - 0.58))
        ctx.renderer.add_table(ctx.slide, table_box, self.headers, self.rows)
        if self.conclusion:
            conclusion_box = Box(box.x + 0.22, box.y + box.h - 0.42, box.w - 0.44, 0.32)
            ctx.renderer.rect(ctx.slide, conclusion_box, ctx.theme.colors.primary_soft, line=ctx.theme.colors.border_light, rounded=True)
            ctx.renderer.text(ctx.slide, conclusion_box.inset(0.10, 0.05), self.conclusion, size=ctx.theme.fonts.caption_size, color=ctx.theme.colors.primary_dark, bold=True, align="center", valign="middle")

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "ComparisonTableComponent":
        rows: list[list[str]] = []
        for row in props.get("rows", []):
            if isinstance(row, Sequence) and not isinstance(row, str):
                rows.append([str(value) for value in row])
        return cls(
            headers=[str(item) for item in props.get("headers", [])],
            rows=rows,
            conclusion=str(props.get("conclusion", "")),
        )


@dataclass
class TimelineComponent:
    items: list[dict[str, Any]] = field(default_factory=list)

    def render(self, ctx: RenderContext, box: Box) -> None:
        r = ctx.renderer
        r.add_card(ctx.slide, box, fill=_style(ctx, "fill", ctx.theme.colors.surface_white), line=_style(ctx, "border", ctx.theme.colors.border_light))
        inner = box.inset(0.30, 0.35, 0.30, 0.30)
        line_y = inner.y + 0.18
        r.line(ctx.slide, inner.x + 0.25, line_y, inner.x + inner.w - 0.25, line_y, color=ctx.theme.colors.border, width=1.0)
        slots = inner.split_cols(max(1, len(self.items)), 0.14)
        for slot, item in zip(slots, self.items):
            r.add_status_timeline_node(
                ctx.slide,
                slot,
                label=str(item.get("label", "")),
                date=str(item.get("date", "")),
                description=str(item.get("description", "")),
                status=str(item.get("status", "normal")),
            )

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "TimelineComponent":
        return cls(items=_items(props.get("items", [])))


@dataclass
class BarChartComponent:
    categories: list[str] = field(default_factory=list)
    series: list[dict[str, Any]] = field(default_factory=list)

    def render(self, ctx: RenderContext, box: Box) -> None:
        r = ctx.renderer
        r.add_card(ctx.slide, box, fill=_style(ctx, "fill", ctx.theme.colors.surface_white), line=_style(ctx, "border", ctx.theme.colors.border_light))
        chart = box.inset(0.34, 0.34, 0.28, 0.34)
        plot = Box(chart.x + 0.45, chart.y + 0.38, chart.w - 0.62, chart.h - 0.86)
        colors = _palette(ctx)
        values = [value for item in self.series for value in _numbers(item.get("values", []))]
        max_value = max(values) if values else 1.0
        caption_size = ctx.theme.fonts.caption_size

        for idx in range(4):
            y = plot.y + idx * plot.h / 3
            r.line(ctx.slide, plot.x, y, plot.x + plot.w, y, ctx.theme.colors.border_light, width=0.5)
        groups = max(1, len(self.categories))
        group_w = plot.w / groups
        series_count = max(1, len(self.series))
        bar_w = min(0.20, group_w / (series_count + 1.2))
        for series_idx, item in enumerate(self.series):
            color = str(item.get("color", "")) or colors[series_idx % len(colors)]
            for idx, value in enumerate(_numbers(item.get("values", []))[:groups]):
                h = (value / max_value) * plot.h if max_value else 0
                x = plot.x + idx * group_w + (group_w - bar_w * series_count) / 2 + series_idx * bar_w
                y = plot.y + plot.h - h
                r.rect(ctx.slide, Box(x, y, bar_w * 0.82, h), color, rounded=True)
                if ctx.style.get("show_values", False):
                    label = f"{value:.1f}".rstrip("0").rstrip(".")
                    r.text(
                        ctx.slide,
                        Box(x - 0.12, max(plot.y - 0.02, y - 0.18), bar_w * 0.82 + 0.24, 0.14),
                        label,
                        size=max(6, caption_size - 1),
                        color=ctx.theme.colors.text_secondary,
                        align="center",
                    )
        for idx, label in enumerate(self.categories):
            x = plot.x + idx * group_w + group_w / 2
            r.text(ctx.slide, Box(x - 0.28, plot.y + plot.h + 0.12, 0.56, 0.12), label, size=caption_size, color=ctx.theme.colors.text_secondary, align="center")
        cursor_x = chart.x + chart.w - min(2.8, len(self.series) * 1.20)
        for idx, item in enumerate(self.series[:3]):
            color = str(item.get("color", "")) or colors[idx % len(colors)]
            r.rect(ctx.slide, Box(cursor_x, chart.y + 0.06, 0.13, 0.04), color, rounded=True)
            r.text(ctx.slide, Box(cursor_x + 0.18, chart.y, 0.92, 0.16), str(item.get("name", "")), size=caption_size, color=ctx.theme.colors.text_secondary)
            cursor_x += 1.10

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "BarChartComponent":
        return cls(categories=[str(item) for item in props.get("categories", [])], series=_items(props.get("series", [])))


@dataclass
class PieChartComponent:
    labels: list[str] = field(default_factory=list)
    values: list[float] = field(default_factory=list)
    donut: bool = False

    def render(self, ctx: RenderContext, box: Box) -> None:
        r = ctx.renderer
        r.add_card(ctx.slide, box, fill=_style(ctx, "fill", ctx.theme.colors.surface_white), line=_style(ctx, "border", ctx.theme.colors.border_light))
        inner = box.inset(0.22, 0.22, 0.22, 0.22)
        data = CategoryChartData()
        labels = self.labels or [f"Item {idx + 1}" for idx in range(len(self.values))]
        values = self.values or [1.0]
        data.categories = labels[: len(values)]
        data.add_series("Share", values)
        chart_type = XL_CHART_TYPE.DOUGHNUT if self.donut else XL_CHART_TYPE.PIE
        legend_w = min(1.05, inner.w * 0.38) if inner.w >= 2.1 else 0.0
        chart_box = Box(inner.x, inner.y + 0.05, inner.w - legend_w - 0.08, inner.h - 0.10)
        chart_shape = ctx.slide.shapes.add_chart(chart_type, Inches(chart_box.x), Inches(chart_box.y), Inches(chart_box.w), Inches(chart_box.h), data)
        chart = chart_shape.chart
        chart.has_legend = False
        chart.has_title = False
        colors = _palette(ctx)
        try:
            for idx, point in enumerate(chart.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = r.rgb(colors[idx % len(colors)])
        except (AttributeError, IndexError):
            pass
        if legend_w:
            total = sum(values) or 1.0
            rows = Box(inner.x + inner.w - legend_w, inner.y + 0.35, legend_w, inner.h - 0.70).split_rows(max(1, len(labels)), 0.04)
            for idx, (row, label) in enumerate(zip(rows, labels)):
                color = colors[idx % len(colors)]
                r.rect(ctx.slide, Box(row.x, row.y + 0.04, 0.10, 0.08), color, rounded=True)
                pct = values[idx] / total * 100 if idx < len(values) else 0
                r.text(ctx.slide, Box(row.x + 0.16, row.y, row.w - 0.16, 0.12), str(label), size=ctx.theme.fonts.caption_size, color=ctx.theme.colors.text_primary, bold=True)
                r.text(ctx.slide, Box(row.x + 0.16, row.y + 0.13, row.w - 0.16, 0.10), f"{pct:.0f}%", size=max(6, ctx.theme.fonts.caption_size - 1), color=ctx.theme.colors.text_tertiary)

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "PieChartComponent":
        values = _numbers(props.get("values", props.get("data", [])))
        labels = [str(item) for item in props.get("labels", props.get("categories", []))]
        return cls(labels=labels, values=values, donut=bool(props.get("donut", False)))


@dataclass
class ProgressBarsComponent:
    items: list[dict[str, Any]] = field(default_factory=list)

    def render(self, ctx: RenderContext, box: Box) -> None:
        r = ctx.renderer
        r.add_card(ctx.slide, box, fill=_style(ctx, "fill", ctx.theme.colors.surface_white), line=_style(ctx, "border", ctx.theme.colors.border_light))
        rows = box.inset(0.30, 0.28).split_rows(max(1, len(self.items)), 0.08)
        for row, item in zip(rows, self.items):
            label = str(item.get("label", ""))
            value = max(0.0, min(100.0, float(item.get("value", 0))))
            color = str(item.get("color", "")) or ctx.theme.colors.primary
            r.text(ctx.slide, Box(row.x, row.y + 0.02, 1.35, 0.18), label, size=ctx.theme.fonts.caption_size, color=ctx.theme.colors.text_primary, bold=True)
            track = Box(row.x + 1.50, row.y + 0.08, row.w - 2.10, 0.08)
            r.rect(ctx.slide, track, ctx.theme.colors.gray_100, rounded=True)
            r.rect(ctx.slide, Box(track.x, track.y, track.w * value / 100, track.h), color, rounded=True)
            r.text(ctx.slide, Box(row.x + row.w - 0.48, row.y + 0.02, 0.48, 0.16), f"{value:.0f}%", size=8, color=ctx.theme.colors.text_secondary, align="right")

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "ProgressBarsComponent":
        return cls(items=_items(props.get("items", [])))


@dataclass
class ModelDiagramComponent:
    kind: str = "generic"
    title: str = ""
    labels: dict[str, Any] = field(default_factory=dict)

    def render(self, ctx: RenderContext, box: Box) -> None:
        kind = self.kind.lower().replace("-", "_")
        if kind in {"fbm", "fbm_net", "bias_mitigation"}:
            self._render_fbm(ctx, box)
        elif kind in {"protogate", "proto_gate", "multimodal_gate"}:
            self._render_protogate(ctx, box)
        else:
            self._render_generic(ctx, box)

    def _label(self, key: str, default: str) -> str:
        value = self.labels.get(key, default)
        return str(value)

    def _list(self, key: str, default: list[str]) -> list[str]:
        value = self.labels.get(key, default)
        if not isinstance(value, list):
            return default
        return [str(item) for item in value]

    def _node(self, ctx: RenderContext, box: Box, title: str, note: str = "", *, fill: str | None = None, accent: bool = False) -> None:
        r = ctx.renderer
        fill_color = fill or (ctx.theme.colors.primary_soft if accent else ctx.theme.colors.surface_white)
        r.add_card(ctx.slide, box, fill=fill_color, line=ctx.theme.colors.border_light, shadow=False)
        if box.h < 0.32:
            r.text(ctx.slide, box.inset(0.06, 0.02), title, size=max(7, ctx.theme.fonts.caption_size - 1), color=ctx.theme.colors.text_primary, bold=True, align="center", valign="middle")
            return
        r.text(ctx.slide, Box(box.x + 0.12, box.y + 0.12, box.w - 0.24, 0.18), title, size=ctx.theme.fonts.caption_size, color=ctx.theme.colors.text_primary, bold=True)
        if note:
            r.text(ctx.slide, Box(box.x + 0.12, box.y + 0.40, box.w - 0.24, max(0.10, box.h - 0.50)), note, size=8, color=ctx.theme.colors.text_secondary)

    def _arrow(self, ctx: RenderContext, x1: float, y1: float, x2: float, y2: float, *, color: str | None = None) -> None:
        r = ctx.renderer
        line_color = color or ctx.theme.colors.primary
        r.line(ctx.slide, x1, y1, x2, y2, line_color, width=0.8)
        r.text(ctx.slide, Box(x2 - 0.06, y2 - 0.08, 0.12, 0.16), ">", size=8, color=line_color, bold=True, align="center", valign="middle")

    def _header(self, ctx: RenderContext, inner: Box) -> float:
        if not self.title:
            return inner.y
        ctx.renderer.text(ctx.slide, Box(inner.x, inner.y, inner.w, 0.22), self.title, size=ctx.theme.fonts.body_size, color=ctx.theme.colors.text_primary, bold=True)
        return inner.y + 0.34

    def _render_fbm(self, ctx: RenderContext, box: Box) -> None:
        r = ctx.renderer
        r.add_card(ctx.slide, box, fill=_style(ctx, "fill", ctx.theme.colors.surface_white), line=_style(ctx, "border", ctx.theme.colors.border_light))
        inner = box.inset(0.26, 0.22)
        y = self._header(ctx, inner)
        note = self._label("note", "")
        note_space = 0.44 if note else 0.04
        diagram_bottom = inner.y + inner.h - note_space
        diagram_h = max(1.20, diagram_bottom - y)
        sources = self._list("sources", ["Source A", "Source B", "Source C"])

        source_area = Box(inner.x, y + 0.08, 1.58, max(0.78, diagram_h - 0.16))
        source_rows = source_area.split_rows(len(sources), 0.08)
        for row, source in zip(source_rows, sources):
            self._node(ctx, row, source, "", fill=ctx.theme.colors.gray_50)

        node_h = min(0.72, max(0.56, diagram_h * 0.46))
        mid_y = y + diagram_h / 2
        main_y = mid_y - node_h / 2
        top_y = max(y + 0.04, mid_y - node_h - 0.12)
        bottom_y = min(diagram_bottom - node_h - 0.02, mid_y + 0.12)
        output_x = inner.x + 9.18

        backbone = Box(inner.x + 2.05, main_y, 1.88, node_h)
        feature = Box(inner.x + 4.42, main_y, 1.78, node_h)
        disease = Box(inner.x + 6.70, top_y, 1.95, node_h)
        bias = Box(inner.x + 6.70, bottom_y, 1.95, node_h)
        output = Box(output_x, main_y, max(1.40, inner.x + inner.w - output_x), node_h)

        self._node(ctx, backbone, self._label("backbone", "Shared backbone"), self._label("backbone_note", "Feature extractor"), accent=True)
        self._node(ctx, feature, self._label("feature", "Fused feature"), self._label("feature_note", "Source-invariant representation"))
        self._node(ctx, disease, self._label("disease_head", "Disease head"), self._label("disease_note", "Classification loss"))
        self._node(ctx, bias, self._label("bias_head", "Bias mitigation"), self._label("bias_note", "Source bias constraint"), fill=ctx.theme.colors.accent_soft)
        self._node(ctx, output, self._label("output", "Diagnosis"), self._label("output_note", "Robust prediction"), accent=True)

        mid_y = backbone.y + backbone.h / 2
        self._arrow(ctx, source_area.x + source_area.w + 0.12, mid_y, backbone.x - 0.10, mid_y)
        self._arrow(ctx, backbone.x + backbone.w + 0.10, mid_y, feature.x - 0.10, mid_y)
        self._arrow(ctx, feature.x + feature.w + 0.10, mid_y, disease.x - 0.10, disease.y + disease.h / 2)
        self._arrow(ctx, feature.x + feature.w + 0.10, mid_y, bias.x - 0.10, bias.y + bias.h / 2, color=ctx.theme.colors.accent)
        self._arrow(ctx, disease.x + disease.w + 0.12, disease.y + disease.h / 2, output.x - 0.10, output.y + output.h / 2)

        if note:
            note_box = Box(inner.x, inner.y + inner.h - 0.30, inner.w, 0.24)
            r.rect(ctx.slide, note_box, ctx.theme.colors.primary_soft, line=ctx.theme.colors.border_light, rounded=True)
            r.text(ctx.slide, note_box.inset(0.12, 0.04), note, size=ctx.theme.fonts.caption_size, color=ctx.theme.colors.primary_dark, bold=True, align="center", valign="middle")

    def _render_protogate(self, ctx: RenderContext, box: Box) -> None:
        r = ctx.renderer
        r.add_card(ctx.slide, box, fill=_style(ctx, "fill", ctx.theme.colors.surface_white), line=_style(ctx, "border", ctx.theme.colors.border_light))
        inner = box.inset(0.26, 0.22)
        y = self._header(ctx, inner)
        modalities = self._list("modalities", ["CFP", "OCT", "OCTA"])
        note = self._label("note", "")
        note_space = 0.44 if note else 0.04
        diagram_bottom = inner.y + inner.h - note_space
        diagram_h = max(1.35, diagram_bottom - y)
        node_h = min(0.70, max(0.54, diagram_h * 0.38))
        lane_gap = max(0.12, min(0.30, diagram_h - node_h * 2))
        top_y = y + 0.04
        bottom_y = min(diagram_bottom - node_h - 0.02, top_y + node_h + lane_gap)
        mid_y = top_y + (bottom_y - top_y) / 2
        classifier_x = inner.x + 10.05
        stack_h = min(0.96, max(0.66, diagram_bottom - bottom_y - 0.02))

        prompt = Box(inner.x, top_y, 1.55, node_h)
        text_encoder = Box(inner.x + 2.05, top_y, 1.55, node_h)
        prototypes = Box(inner.x + 4.10, top_y, 1.78, node_h)
        modality_stack = Box(inner.x, bottom_y, 1.55, stack_h)
        visual_encoder = Box(inner.x + 2.05, bottom_y, 1.55, node_h)
        projection = Box(inner.x + 4.10, bottom_y, 1.78, node_h)
        gate = Box(inner.x + 6.38, mid_y, 1.45, node_h)
        fusion = Box(inner.x + 8.25, mid_y, 1.42, node_h)
        classifier = Box(classifier_x, mid_y, max(1.20, inner.x + inner.w - classifier_x), node_h)

        self._node(ctx, prompt, self._label("prompt", "Text prompts"), self._label("prompt_note", "Disease semantics"))
        self._node(ctx, text_encoder, self._label("text_encoder", "Text encoder"), self._label("text_note", "Semantic vectors"), accent=True)
        self._node(ctx, prototypes, self._label("prototype", "Prototype bank"), self._label("prototype_note", "Class anchors"), fill=ctx.theme.colors.accent_soft)
        rows = modality_stack.split_rows(len(modalities), 0.06)
        for row, modality in zip(rows, modalities):
            self._node(ctx, row, modality, "", fill=ctx.theme.colors.gray_50)
        self._node(ctx, visual_encoder, self._label("visual_encoder", "Modality encoders"), self._label("visual_note", "Visual features"), accent=True)
        self._node(ctx, projection, self._label("projection", "Semantic alignment"), self._label("projection_note", "Shared space"))
        self._node(ctx, gate, self._label("gate", "Gate network"), self._label("gate_note", "Adaptive weights"), fill=ctx.theme.colors.accent_soft)
        self._node(ctx, fusion, self._label("fusion", "Weighted fusion"), self._label("fusion_note", "Evidence aggregation"))
        self._node(ctx, classifier, self._label("classifier", "Diagnosis"), self._label("classifier_note", "Known / unknown inference"), accent=True)

        self._arrow(ctx, prompt.x + prompt.w + 0.10, prompt.y + prompt.h / 2, text_encoder.x - 0.10, text_encoder.y + text_encoder.h / 2)
        self._arrow(ctx, text_encoder.x + text_encoder.w + 0.10, text_encoder.y + text_encoder.h / 2, prototypes.x - 0.10, prototypes.y + prototypes.h / 2)
        self._arrow(ctx, modality_stack.x + modality_stack.w + 0.10, visual_encoder.y + visual_encoder.h / 2, visual_encoder.x - 0.10, visual_encoder.y + visual_encoder.h / 2)
        self._arrow(ctx, visual_encoder.x + visual_encoder.w + 0.10, visual_encoder.y + visual_encoder.h / 2, projection.x - 0.10, projection.y + projection.h / 2)
        self._arrow(ctx, prototypes.x + prototypes.w + 0.10, prototypes.y + prototypes.h / 2, gate.x - 0.10, gate.y + gate.h / 2, color=ctx.theme.colors.accent)
        self._arrow(ctx, projection.x + projection.w + 0.10, projection.y + projection.h / 2, gate.x - 0.10, gate.y + gate.h / 2)
        self._arrow(ctx, gate.x + gate.w + 0.10, gate.y + gate.h / 2, fusion.x - 0.10, fusion.y + fusion.h / 2)
        self._arrow(ctx, fusion.x + fusion.w + 0.10, fusion.y + fusion.h / 2, classifier.x - 0.10, classifier.y + classifier.h / 2)

        if note:
            note_box = Box(inner.x, inner.y + inner.h - 0.30, inner.w, 0.24)
            r.rect(ctx.slide, note_box, ctx.theme.colors.primary_soft, line=ctx.theme.colors.border_light, rounded=True)
            r.text(ctx.slide, note_box.inset(0.12, 0.04), note, size=ctx.theme.fonts.caption_size, color=ctx.theme.colors.primary_dark, bold=True, align="center", valign="middle")

    def _render_generic(self, ctx: RenderContext, box: Box) -> None:
        r = ctx.renderer
        r.add_card(ctx.slide, box, fill=_style(ctx, "fill", ctx.theme.colors.surface_white), line=_style(ctx, "border", ctx.theme.colors.border_light))
        stages = self._list("stages", ["Input", "Encoder", "Fusion", "Prediction"])
        inner = box.inset(0.26, 0.24)
        y = self._header(ctx, inner)
        cells = Box(inner.x, y + 0.15, inner.w, max(0.60, inner.y + inner.h - y - 0.25)).split_cols(len(stages), 0.18)
        for idx, (cell, stage) in enumerate(zip(cells, stages)):
            self._node(ctx, cell, stage, "", accent=idx in {0, len(stages) - 1})
            if idx < len(cells) - 1:
                self._arrow(ctx, cell.x + cell.w + 0.04, cell.y + cell.h / 2, cells[idx + 1].x - 0.06, cell.y + cell.h / 2)

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "ModelDiagramComponent":
        labels = dict(props)
        kind = str(labels.pop("kind", labels.pop("variant", "generic")))
        title = str(labels.pop("title", ""))
        return cls(kind=kind, title=title, labels=labels)


@dataclass
class ProcessFlowComponent:
    steps: list[dict[str, Any]] = field(default_factory=list)

    def render(self, ctx: RenderContext, box: Box) -> None:
        steps = self.steps[:6]
        if not steps:
            return
        r = ctx.renderer
        r.add_card(ctx.slide, box, fill=_style(ctx, "fill", ctx.theme.colors.surface_white), line=_style(ctx, "border", ctx.theme.colors.border_light))
        compact = bool(ctx.style.get("compact", box.h < 1.85))
        inner = box.inset(0.22, 0.24 if compact else 0.34, 0.22, 0.22 if compact else 0.28)
        cells = inner.split_cols(len(steps), _style_num(ctx, "gap", 0.16))
        max_card_h = _style_num(ctx, "card_height", 1.22 if compact else 1.55)
        output_label = str(ctx.style.get("output_label", "Output"))
        for idx, (cell, step) in enumerate(zip(cells, steps), start=1):
            card_h = min(cell.h, max_card_h)
            card_y = cell.y + max(0, (cell.h - card_h) / 2)
            step_box = Box(cell.x, card_y, cell.w, card_h)
            r.add_process_step_card(
                ctx.slide,
                step_box,
                index=idx,
                title=str(step.get("title", step.get("label", ""))),
                description=str(step.get("description", "")),
                output=str(step.get("output", "")),
                compact=compact,
                output_label=output_label,
            )
            if idx < len(cells):
                x = cell.x + cell.w + 0.03
                y = card_y + card_h / 2
                r.line(ctx.slide, x, y, x + 0.10, y, ctx.theme.colors.primary, width=0.8)

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "ProcessFlowComponent":
        return cls(steps=_items(props.get("steps", props.get("items", []))))


@dataclass
class RoadmapComponent:
    items: list[dict[str, Any]] = field(default_factory=list)

    def render(self, ctx: RenderContext, box: Box) -> None:
        r = ctx.renderer
        r.add_card(ctx.slide, box, fill=_style(ctx, "fill", ctx.theme.colors.surface_white), line=_style(ctx, "border", ctx.theme.colors.border_light))
        inner = box.inset(0.34, 0.30)
        rows = inner.split_rows(max(1, len(self.items)), 0.10)
        colors = _palette(ctx)
        for idx, (row, item) in enumerate(zip(rows, self.items)):
            color = str(item.get("color", "")) or colors[idx % len(colors)]
            r.text(ctx.slide, Box(row.x, row.y + 0.02, 1.0, 0.18), str(item.get("label", "")), size=8, color=ctx.theme.colors.text_primary, bold=True)
            track = Box(row.x + 1.10, row.y + 0.08, row.w - 1.55, 0.07)
            r.rect(ctx.slide, track, ctx.theme.colors.gray_100, rounded=True)
            start = max(0.0, min(1.0, float(item.get("start", 0))))
            end = max(start + 0.05, min(1.0, float(item.get("end", 1))))
            r.rect(ctx.slide, Box(track.x + track.w * start, track.y, track.w * (end - start), track.h), color, rounded=True)
            r.text(ctx.slide, Box(row.x + row.w - 0.40, row.y + 0.01, 0.40, 0.16), str(item.get("period", "")), size=7, color=ctx.theme.colors.text_tertiary, align="right")

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "RoadmapComponent":
        return cls(items=_items(props.get("items", [])))


@dataclass
class IconComponent:
    name: str = ""
    label: str = ""
    source: str = ""
    description: str = ""
    src: str = ""
    color: str = ""
    size: int | None = None
    width: int | None = None
    height: int | None = None
    rotate: str | int | None = None
    flip: str = ""
    stroke_width: float | None = None
    opacity: float = 1.0

    def _render_mark(self, ctx: RenderContext, icon_box: Box, accent: str, fallback: str, *, size: int = 9) -> None:
        if self.src and ctx.renderer.picture(ctx.slide, self.src, icon_box.inset(0.08, 0.08), fit="contain") is not None:
            return
        request = ctx.renderer.icon_registry.create_request(
            self.name,
            color=accent,
            size=self.size or self.width or self.height or 128,
            width=self.width,
            height=self.height,
            rotate=self.rotate,
            flip=self.flip,
            stroke_width=self.stroke_width,
        )
        if request is not None and ctx.renderer.icon_picture(ctx.slide, request, icon_box.inset(0.08, 0.08), opacity=self.opacity) is not None:
            return
        ctx.renderer.text(ctx.slide, icon_box.inset(0.03, 0.02), fallback, size=size, color=accent, bold=True, align="center", valign="middle")

    def render(self, ctx: RenderContext, box: Box) -> None:
        r = ctx.renderer
        fill = _style(ctx, "fill", ctx.theme.colors.primary_soft)
        accent = _color(ctx, self.color or ctx.style.get("accent"), ctx.theme.colors.primary)
        r.add_card(ctx.slide, box, fill=fill, line=_style(ctx, "border", ctx.theme.colors.border_light), shadow=bool(ctx.style.get("shadow", False)))
        icon_text = self.label or self.name.split(":")[-1][:2].upper()
        if box.h < 1.15:
            center_y = box.y + box.h / 2 - (0.10 if self.name else 0)
            icon_box = Box(box.x + box.w / 2 - 0.20, center_y - 0.20, 0.40, 0.40)
            r.circle(ctx.slide, icon_box, "FFFFFF", line=ctx.theme.colors.border_light)
            self._render_mark(ctx, icon_box, accent, icon_text, size=8)
            if self.name:
                r.text(ctx.slide, Box(box.x + 0.10, center_y + 0.30, box.w - 0.20, 0.14), self.name, size=7, color=ctx.theme.colors.text_secondary, align="center")
            return

        inner = box.inset(0.22, 0.18)
        icon_box = Box(inner.x, inner.y + 0.02, 0.48, 0.48)
        r.circle(ctx.slide, icon_box, "FFFFFF", line=ctx.theme.colors.border_light)
        self._render_mark(ctx, icon_box, accent, icon_text, size=9)
        title = self.source or self.name.split(":")[0].title() or "Icon"
        r.text(ctx.slide, Box(inner.x + 0.62, inner.y + 0.04, inner.w - 0.62, 0.20), title, size=ctx.theme.fonts.caption_size, color=ctx.theme.colors.text_primary, bold=True)
        r.text(ctx.slide, Box(inner.x + 0.62, inner.y + 0.28, inner.w - 0.62, 0.14), self.name, size=7, color=ctx.theme.colors.text_tertiary)
        if self.description:
            r.text(ctx.slide, Box(inner.x, inner.y + 0.72, inner.w, inner.h - 0.72), self.description, size=8, color=ctx.theme.colors.text_secondary)

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "IconComponent":
        return cls(
            name=str(props.get("name", "")),
            label=str(props.get("label", "")),
            source=str(props.get("source", "")),
            description=str(props.get("description", "")),
            src=str(props.get("src", "")),
            color=str(props.get("color", "")),
            size=int(props["size"]) if props.get("size") is not None else None,
            width=int(props["width"]) if props.get("width") is not None else None,
            height=int(props["height"]) if props.get("height") is not None else None,
            rotate=props.get("rotate"),
            flip=str(props.get("flip", "")),
            stroke_width=float(props["stroke_width"]) if props.get("stroke_width") is not None else (float(props["strokeWidth"]) if props.get("strokeWidth") is not None else None),
            opacity=float(props.get("opacity", 1.0)),
        )


@dataclass
class ImageComponent:
    src: str = ""
    fit: str = "contain"

    def render(self, ctx: RenderContext, box: Box) -> None:
        path = Path(self.src)
        if not path.exists():
            ctx.renderer.add_card(ctx.slide, box, fill=ctx.theme.colors.gray_50, line=ctx.theme.colors.border_light)
            ctx.renderer.text(ctx.slide, box.inset(0.18, 0.16), f"Missing image: {self.src}", size=8, color=ctx.theme.colors.text_tertiary, align="center", valign="middle")
            return
        ctx.renderer.picture(ctx.slide, path, box, fit=self.fit)

    @classmethod
    def from_props(cls, props: Mapping[str, Any]) -> "ImageComponent":
        return cls(src=str(props.get("src", "")), fit=str(props.get("fit", "contain")))
